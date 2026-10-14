"""Stage 1a helper — Brain 1's files on disk -> tmp/brain1/<user>/*.txt

Reads every `type: filesystem, enabled: true` source in the config (or one
ad-hoc --path), converts each supported file to plain text with a metadata
header, and skips files whose mtime is unchanged since the last run
(`.source-index.json`). Output name: <sanitised_basename>__<category>__fs.txt

Usage:
  python capture_fs.py <user>
  python capture_fs.py <user> --path ~/Desktop/note.pdf --category notes
"""

from __future__ import annotations

import argparse
import csv
import io
import re
from pathlib import Path

from _common import load_config, load_json, now_iso, sanitise, save_json, sha256_text, sources_dir

SUPPORTED = {'.md', '.txt', '.pdf', '.csv', '.xlsx'}


def convert(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in ('.md', '.txt'):
        return path.read_text(encoding='utf-8', errors='replace')
    if suffix == '.pdf':
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return '\n\n'.join(f'[page {i + 1}]\n{(page.extract_text() or "").strip()}' for i, page in enumerate(reader.pages))
    if suffix == '.csv':
        with path.open(newline='', encoding='utf-8', errors='replace') as fh:
            return _rows_to_text(list(csv.reader(fh)))
    if suffix == '.xlsx':
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        chunks = []
        for ws in wb.worksheets:
            rows = [['' if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
            chunks.append(f'[sheet {ws.title}]\n{_rows_to_text(rows)}')
        return '\n\n'.join(chunks)
    raise ValueError(f'unsupported file type: {suffix}')


def _rows_to_text(rows: list[list[str]]) -> str:
    out = io.StringIO()
    csv.writer(out, delimiter='|').writerows(rows)
    return out.getvalue()


def iter_files(root: Path, recursive: bool, extensions: set[str]):
    if root.is_file():
        yield root
        return
    pattern = '**/*' if recursive else '*'
    for p in sorted(root.glob(pattern)):
        if p.is_file() and p.suffix.lower() in extensions and not p.name.startswith('.'):
            yield p


def capture(user: str, entries: list[dict]) -> dict:
    out_dir = sources_dir(user)
    out_dir.mkdir(parents=True, exist_ok=True)
    index_path = out_dir / '.source-index.json'
    index = load_json(index_path, {})
    stats = {'written': 0, 'unchanged': 0, 'failed': []}

    for entry in entries:
        root = Path(entry['path']).expanduser()
        if not root.exists():
            stats['failed'].append(f"{entry.get('name', root)}: path not found {root}")
            continue
        exts = {e.lower() for e in entry.get('extensions', SUPPORTED)} & SUPPORTED
        category = sanitise(entry.get('category', 'notes'), 30)
        for path in iter_files(root, entry.get('recursive', True), exts):
            key = f'fs:{path.resolve()}'
            version = str(path.stat().st_mtime_ns)
            out_name = f'{sanitise(path.stem)}__{category}__fs.txt'
            taken_by = {v['file']: k for k, v in index.items()}
            if taken_by.get(out_name, key) != key:  # same name from a different folder
                out_name = f'{sanitise(path.stem)}-{sha256_text(str(path.resolve()))[:6]}__{category}__fs.txt'
            cached = index.get(key)
            if cached and cached['version'] == version and (out_dir / cached['file']).exists():
                stats['unchanged'] += 1
                continue
            try:
                body = convert(path)
            except Exception as exc:  # one bad file must not stop the run
                stats['failed'].append(f'{path}: {exc}')
                continue
            if not re.sub(r'\[(page|sheet) [^\]]*\]', '', body).strip():
                stats['failed'].append(f'{path}: no extractable text (scanned image or empty file?)')
                continue
            header = (
                f'source: filesystem\noriginal_path: {path}\ncategory: {category}\n'
                f'modified_ns: {version}\ncaptured_at: {now_iso()}\n---\n'
            )
            (out_dir / out_name).write_text(header + body, encoding='utf-8')
            index[key] = {'file': out_name, 'version': version, 'source': 'filesystem', 'category': category}
            stats['written'] += 1

    save_json(index_path, index)
    write_index_md(user)
    return stats


def write_index_md(user: str) -> None:
    out_dir = sources_dir(user)
    index = load_json(out_dir / '.source-index.json', {})
    lines = [f'# Brain 1 sources — {user}', '', '| File | Source | Category |', '|---|---|---|']
    for key, item in sorted(index.items(), key=lambda kv: kv[1]['file']):
        lines.append(f"| {item['file']} | {item['source']} | {item.get('category', '')} |")
    (out_dir / 'index.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('user')
    parser.add_argument('--path', help='ad-hoc: capture one file or folder instead of the config')
    parser.add_argument('--category', default='notes')
    args = parser.parse_args()

    if args.path:
        entries = [{'name': 'ad-hoc', 'path': args.path, 'category': args.category}]
    else:
        cfg = load_config(args.user)
        entries = [s for s in cfg['sources'] if s.get('type') == 'filesystem' and s.get('enabled')]
    stats = capture(args.user, entries)
    print(f"filesystem: {stats['written']} written, {stats['unchanged']} unchanged, {len(stats['failed'])} failed")
    for f in stats['failed']:
        print(f'  FAILED {f}')


if __name__ == '__main__':
    main()
